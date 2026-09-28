# -*- coding: utf-8 -*-
"""Genera la 'Solicitud de Móvil Municipal' en Word, replicando el formulario original."""
import sys
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

COLS = [184, 10, 99, 28, 94, 26, 99]  # puntos; suman 540 (ancho útil de una hoja carta)
FONT = "Arial"
ESCALA = 0.96  # escala vertical de todas las filas (deja margen para que quepa en una hoja)

_ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER}
_VALIGN = {"top": WD_ALIGN_VERTICAL.TOP, "center": WD_ALIGN_VERTICAL.CENTER}


def _bordes(cell, top=False, left=False, bottom=False, right=False):
    tcPr = cell._tc.get_or_add_tcPr()
    actual = tcPr.find(qn("w:tcBorders"))
    aristas = {"top": top, "left": left, "bottom": bottom, "right": right}
    if actual is not None:
        for e in aristas:
            if actual.find(qn("w:" + e)) is not None:
                aristas[e] = True
        tcPr.remove(actual)
    b = OxmlElement("w:tcBorders")
    for e in ("top", "left", "bottom", "right"):
        if aristas[e]:
            el = OxmlElement("w:" + e)
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "8")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
            b.append(el)
    tcPr.insert_element_before(b, "w:shd", "w:noWrap", "w:tcMar", "w:textDirection",
                               "w:tcFitText", "w:vAlign", "w:hideMark")


def _caja(cell):
    _bordes(cell, True, True, True, True)


def _texto(cell, texto, size=10, bold=False, italic=False, underline=False,
           align="center", valign="center", indent=0):
    cell.vertical_alignment = _VALIGN[valign]
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.alignment = _ALIGN[align]
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    if indent:
        pf.left_indent = Pt(indent)
    partes = (texto or "").split("\n")
    for i, parte in enumerate(partes):
        run = p.add_run(parte)
        run.font.name = FONT
        run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.underline = underline
        if i < len(partes) - 1:
            run.add_break()


def _fila(tabla, alto, exacto=False):
    fila = tabla.add_row()
    fila.height = Pt(alto * ESCALA)
    fila.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if exacto else WD_ROW_HEIGHT_RULE.AT_LEAST
    for i, c in enumerate(fila.cells):
        c.width = Pt(COLS[i])
    return fila


def _preparar_tabla(tabla):
    tabla.autofit = False
    tblPr = tabla._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(sum(COLS) * 20))
    mar = OxmlElement("w:tblCellMar")
    for lado, v in (("top", 0), ("left", 40), ("bottom", 0), ("right", 40)):
        el = OxmlElement("w:" + lado)
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def build_form(d, ruta_salida):
    """d: dict con las claves creado_fecha, creado_hora, unidad, funcionario, cargo,
    srv_fecha, hora_inicio, hora_termino, coordinador, conductor, lugar_inicio,
    lugar_termino, cometido."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(612), Pt(792)
    sec.left_margin = sec.right_margin = Pt(36)
    sec.top_margin = sec.bottom_margin = Pt(28)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)

    t = doc.add_table(rows=0, cols=7)
    _preparar_tabla(t)

    def espacio(h):
        _fila(t, h, exacto=True)

    def campo(etiqueta, valor, ancho="completo", size=10, bold=False, italic=False,
              alto=15, valign="center"):
        f = _fila(t, alto)
        c = f.cells
        _texto(c[0], etiqueta, 8, True, True, align="left", valign=valign)
        _texto(c[1], ":", 10, True, align="center", valign=valign)
        caja = c[2].merge(c[6]) if ancho == "completo" else c[2]
        _texto(caja, valor, size, bold, italic, align="center", valign="center")
        _caja(caja)
        return f

    # Título
    f = _fila(t, 18)
    titulo = f.cells[0].merge(f.cells[6])
    _texto(titulo, "SOLICITUD DE MÓVIL MUNICIPAL", 11, True, True, True, "center", indent=80)
    espacio(12)

    # Fecha y hora de creación
    for etiqueta, valor in (("FECHA", d.get("creado_fecha", "")), ("HORA", d.get("creado_hora", ""))):
        f = _fila(t, 15)
        c = f.cells
        _texto(c[4], etiqueta, 10, True, True, align="left")
        _texto(c[5], ":", 10, True, align="center")
        _texto(c[6], valor, 9, align="center")
        _caja(c[6])
        espacio(8)

    # 1.- Antecedentes del requiriente
    f = _fila(t, 14)
    _texto(f.cells[0].merge(f.cells[6]), "1.- ANTECEDENTES DEL REQUIRIENTE", 10, True, True, True, "left")
    espacio(8)
    campo("UNIDAD", d.get("unidad", ""))
    espacio(9)
    campo("NOMBRE FUNCIONARIO", d.get("funcionario", ""))
    espacio(9)
    campo("CARGO", d.get("cargo", ""))
    espacio(12)

    # 2.- Antecedentes del servicio
    f = _fila(t, 14)
    _texto(f.cells[0].merge(f.cells[6]), "2.- ANTECEDENTES DEL SERVICIO", 10, True, True, True, "left")
    espacio(8)
    campo("FECHA", d.get("srv_fecha", ""), ancho="chico", bold=True, italic=True)
    espacio(9)
    campo("HORA INICIO", d.get("hora_inicio", ""), ancho="chico", bold=True, italic=True)
    espacio(9)
    campo("HORA TERMINO", d.get("hora_termino", ""), ancho="chico", bold=True, italic=True)
    espacio(9)

    # Jornada de colación (siempre en blanco)
    for i, turno in enumerate(("(MAÑANA)", "(TARDE)", "(NOCHE)")):
        f = _fila(t, 14)
        c = f.cells
        if i == 0:
            _texto(c[0], "* JORNADA DE COLACIÓN", 8, True, True, align="left")
            _texto(c[1], ":", 10, True, align="center")
        _caja(c[2])
        _texto(c[3], "A", 10, True, align="center")
        _caja(c[3])
        _caja(c[4])
        _texto(c[5].merge(c[6]), turno, 10, align="left", indent=4)
    espacio(10)

    campo("COORDINADOR RESPONSABLE\nDEL MÓVIL", d.get("coordinador", ""), bold=True, italic=True, alto=28)
    espacio(9)
    campo("CONDUCTOR TITULAR", d.get("conductor", ""), bold=True, italic=True)
    espacio(9)
    campo("LUGAR DE INICIO SERVICIO", d.get("lugar_inicio", ""), bold=True, italic=True)
    espacio(9)
    campo("LUGAR DE TÉRMINO SERVICIO", d.get("lugar_termino", ""), bold=True, italic=True)
    espacio(9)
    campo("COMETIDO A REALIZAR", d.get("cometido", ""), bold=True, italic=True, alto=78, valign="top")
    espacio(16)

    # Uso exclusivo de la sección de transporte (en blanco)
    caja_filas = []
    f = _fila(t, 17)
    _texto(f.cells[0].merge(f.cells[6]), "USO EXCLUSIVO DE LA SECCIÓN DE TRANSPORTE", 10, True, False, True, "left", indent=6)
    caja_filas.append(f)
    caja_filas.append(_fila(t, 6, exacto=True))

    f = _fila(t, 15)
    c = f.cells
    _texto(c[0], "Nº DE SOLICITUD", 8, True, True, align="left")
    _texto(c[1], ":", 10, True, align="center")
    _caja(c[2])
    caja_filas.append(f)
    caja_filas.append(_fila(t, 5, exacto=True))

    f = _fila(t, 15)
    c = f.cells
    _texto(c[0], "SOLICITUD", 8, True, True, align="left")
    _texto(c[1], ":", 10, True, align="center")
    _texto(c[2], "APROBACIÓN", 8, True, True, align="left")
    _caja(c[3])
    _texto(c[4], "RECHAZO", 8, True, True, align="left", indent=4)
    _caja(c[5])
    caja_filas.append(f)
    caja_filas.append(_fila(t, 5, exacto=True))

    f = _fila(t, 15)
    c = f.cells
    _texto(c[0], "MOTIVO DE RECHAZO", 8, True, True, align="left")
    _texto(c[1], ":", 10, True, align="center")
    _caja(c[2].merge(c[6]))
    caja_filas.append(f)
    caja_filas.append(_fila(t, 16, exacto=True))

    # Borde exterior del recuadro de Transporte
    ultima = len(caja_filas) - 1
    for i, fila in enumerate(caja_filas):
        tcs = fila._tr.tc_lst
        from docx.table import _Cell
        for j, tc in enumerate(tcs):
            cel = _Cell(tc, t)
            _bordes(cel,
                    top=(i == 0),
                    left=(j == 0),
                    bottom=(i == ultima),
                    right=(j == len(tcs) - 1))

    # Párrafo final mínimo (Word exige uno después de una tabla)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)
    r = p.add_run("")
    r.font.size = Pt(1)

    doc.save(ruta_salida)
    return ruta_salida


DEMO = {
    "creado_fecha": "09-09-2026", "creado_hora": "12:10:00",
    "unidad": "DIRECCIÓN DE SERVICIOS A LA COMUNIDAD",
    "funcionario": "ALEJANDRO GALLEGUILLOS RIVERA",
    "cargo": "DIRECTOR DE SERVICIOS A LA COMUNIDAD (S)",
    "srv_fecha": "12-09-2026", "hora_inicio": "8:00:00", "hora_termino": "23:59:00",
    "coordinador": "ALEJANDRO GALLEGUILLOS RIVERA",
    "conductor": "A DESIGNAR POR TRANSPORTE",
    "lugar_inicio": "BODEGA DE DISERCO, CALLE RENGIFO",
    "lugar_termino": "BODEGA DE DISERCO, CALLE RENGIFO",
    "cometido": "MANTENCIÓN AREAS VERDES NO CONCESIONADAS .",
}

if __name__ == "__main__":
    salida = sys.argv[1] if len(sys.argv) > 1 else "demo.docx"
    build_form(DEMO, salida)
    print("OK", salida)
