# -*- coding: utf-8 -*-
"""Genera la 'Solicitud de Móvil Municipal' en Excel a partir de la plantilla (copia del formulario original)."""
import sys
from datetime import datetime, time
from pathlib import Path

from openpyxl import load_workbook

PLANTILLA = Path(__file__).parent / "plantilla_movil.xlsx"
MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE",
         "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]


def _fecha(s):
    return datetime.strptime(s, "%d-%m-%Y")


def _hora(s):
    h, m, seg = (int(x) for x in s.split(":"))
    return time(h, m, seg)


def build_form(d, ruta_salida):
    """d: mismo dict que usaba el formulario en Word (creado_fecha, creado_hora, unidad, funcionario, cargo,
    srv_fecha, hora_inicio, hora_termino, coordinador, conductor, lugar_inicio, lugar_termino, cometido)."""
    wb = load_workbook(PLANTILLA)
    ws = wb.active
    f = _fecha(d["srv_fecha"])
    ws.title = "%02d %s" % (f.day, MESES[f.month - 1])

    ws["I5"] = _fecha(d["creado_fecha"])
    ws["I7"] = _hora(d["creado_hora"])
    ws["E11"] = d["unidad"]
    ws["E13"] = d["funcionario"]
    ws["E15"] = d["cargo"]
    ws["E19"] = f
    ws["E21"] = _hora(d["hora_inicio"])
    ws["E23"] = _hora(d["hora_termino"])
    ws["E29"] = d["coordinador"] or None
    ws["E31"] = d["conductor"]
    ws["E33"] = d["lugar_inicio"]
    ws["E35"] = d["lugar_termino"]
    ws["E37"] = d["cometido"]
    for celda in ("I5", "E19"):
        ws[celda].number_format = "dd-mm-yyyy;@"
    for celda in ("I7", "E21", "E23"):
        ws[celda].number_format = "h:mm:ss;@"

    wb.save(ruta_salida)
    return ruta_salida


DEMO = {
    "creado_fecha": "09-09-2026", "creado_hora": "12:10:00",
    "unidad": "DIRECCIÓN DE SERVICIOS A LA COMUNIDAD",
    "funcionario": "ALEJANDRO GALLEGUILLOS RIVERA",
    "cargo": "DIRECTOR DE SERVICIOS A LA COMUNIDAD (S)",
    "srv_fecha": "12-09-2026", "hora_inicio": "8:00:00", "hora_termino": "23:59:00",
    "coordinador": "ALEJANDRO GALLEGUILLOS RIVERA",
    "conductor": "A DESIGNAR POR TRANSPORTE",
    "lugar_inicio": "BODEGA DE DISERCO, CALLE RENGIFO",
    "lugar_termino": "BODEGA DE DISERCO, CALLE RENGIFO",
    "cometido": "MANTENCIÓN AREAS VERDES NO CONCESIONADAS .",
}

if __name__ == "__main__":
    salida = sys.argv[1] if len(sys.argv) > 1 else "demo.xlsx"
    build_form(DEMO, salida)
    print("OK", salida)
